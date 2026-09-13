import streamlit as st
import pandas as pd
import requests
from datetime import datetime
from urllib.parse import quote
from io import BytesIO

# ============================================================
# NUSTATYMAI
# ============================================================
N8N_BASE_URL = st.secrets.get("N8N_BASE_URL", "")
N8N_API_KEY = st.secrets.get("N8N_API_KEY", "")

ALLOWED_STATUSES = ["Laukia", "Patvirtinta"]
REQUEST_TIMEOUT = 20
NUMERIC_COLUMNS = ["kiekis", "periodiskumas", "ikainis", "suma_be_pvm", "suma_su_pvm"]
# ============================================================

st.set_page_config(page_title="Sumų suvestinė", page_icon="📊", layout="wide")
st.title("📊 Sumų suvestinė — valdymo skydelis")

if not N8N_BASE_URL:
    st.error(
        "Nenustatytas N8N_BASE_URL. Programos nustatymuose (Settings → Secrets) "
        "įrašyk N8N_BASE_URL ir N8N_API_KEY."
    )
    st.stop()


def headers():
    return {"X-Api-Key": N8N_API_KEY}


@st.cache_data(ttl=15, show_spinner=False)
def fetch_data():
    """
    GET {N8N_BASE_URL}/webhook/gauti-duomenis
    Tikimasi n8n Webhook (GET), kuris perziurai.xlsx paverčia JSON ir
    grąžina per "Respond to Webhook" mazgą. Turinys — visų eilučių
    sąrašas (įskaitant "VISO:" eilutę), pvz.:
        [{"klientas": "...", "sutarties_nr": "...", "objektas": "...",
          "suma_be_pvm": 123.45, "suma_su_pvm": 149.37,
          "busena": "Laukia", "saltinio_failas": "...", ...}, ...]
    """
    resp = requests.get(
        f"{N8N_BASE_URL}/webhook/gauti-duomenis", headers=headers(), timeout=REQUEST_TIMEOUT
    )
    resp.raise_for_status()
    df = pd.DataFrame(resp.json())
    return clean_numeric_columns(df)


def clean_numeric_columns(dataframe: pd.DataFrame) -> pd.DataFrame:
    """Priverstinai konvertuoja žinomus skaitinius stulpelius į skaičius,
    nes n8n juos grąžina kaip tekstą, o Excel eksportas tada rašo juos
    kaip tekstines, ne skaitines, celes."""
    df = dataframe.copy()
    for col in NUMERIC_COLUMNS:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    return df


def post_status_updates(rows: list[dict]):
    """
    POST {N8N_BASE_URL}/webhook/atnaujinti-busena
    Body: {"updates": [{"saltinio_failas": "...", "busena": "Patvirtinta"}, ...]}
    n8n pusėje: pagal saltinio_failas surasti atitinkamą eilutę
    perziurai.xlsx faile ir perrašyti busena stulpelį.
    """
    resp = requests.post(
        f"{N8N_BASE_URL}/webhook/atnaujinti-busena",
        headers=headers(),
        json={"updates": rows},
        timeout=REQUEST_TIMEOUT,
    )
    resp.raise_for_status()


def to_excel_bytes(dataframe: pd.DataFrame) -> bytes:
    """
    Konvertuoja DataFrame į .xlsx baitus su TARPINĖMIS SUMOMIS (subtotal):

    - eilutės sugrupuojamos pagal stulpelį "klientas";
    - po kiekvieno kliento eilučių įterpiama tarpinės sumos eilutė su
      tikra Excel formule =SUBTOTAL(9; ...) stulpeliams suma_be_pvm ir
      suma_su_pvm;
    - kliento eilutės paslepiamos į grupę (outline), kairėje Excel pusėje
      atsiranda [+]/[-] mygtukai joms suskleisti/išskleisti;
    - apačioje bendra suma (VISO:), taip pat kaip =SUBTOTAL(9; ...) formulė
      per visą lentelę — SUBTOTAL funkcija automatiškai ignoruoja kitas
      SUBTOTAL eilutes, todėl bendra suma neduosis dvigubai.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    df = dataframe.copy()
    if "klientas" in df.columns:
        df = df[df["klientas"] != "VISO:"].copy()
        df = df.sort_values("klientas", kind="stable").reset_index(drop=True)

    columns = list(df.columns)
    subtotal_cols = [c for c in ("suma_be_pvm", "suma_su_pvm") if c in columns]
    can_group = "klientas" in columns and subtotal_cols

    wb = Workbook()
    ws = wb.active
    ws.title = "Suvestine"
    bold = Font(bold=True)

    ws.append(columns)
    for cell in ws[1]:
        cell.font = bold

    current_row = 2

    if can_group:
        klientas_idx = columns.index("klientas")
        for klientas, group in df.groupby("klientas", sort=False):
            group_start = current_row
            for _, row in group.iterrows():
                ws.append([row[c] for c in columns])
                ws.row_dimensions[current_row].outline_level = 1
                current_row += 1
            group_end = current_row - 1

            subtotal_row = ["" for _ in columns]
            subtotal_row[klientas_idx] = f"Tarpinė suma: {klientas}"
            ws.append(subtotal_row)
            for c in subtotal_cols:
                col_idx = columns.index(c) + 1
                col_letter = get_column_letter(col_idx)
                cell = ws.cell(row=current_row, column=col_idx)
                cell.value = f"=SUBTOTAL(9,{col_letter}{group_start}:{col_letter}{group_end})"
            for cell in ws[current_row]:
                cell.font = bold
            current_row += 1

        grand_total_end = current_row - 1
        grand_row = ["" for _ in columns]
        grand_row[klientas_idx] = "VISO:"
        ws.append(grand_row)
        for c in subtotal_cols:
            col_idx = columns.index(c) + 1
            col_letter = get_column_letter(col_idx)
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = f"=SUBTOTAL(9,{col_letter}2:{col_letter}{grand_total_end})"
        for cell in ws[current_row]:
            cell.font = bold
    else:
        for _, row in df.iterrows():
            ws.append([row[c] for c in columns])

    ws.sheet_properties.outlinePr.summaryBelow = True

    for col_idx, col_name in enumerate(columns, start=1):
        max_len = max([len(str(col_name))] + [len(str(v)) for v in df[col_name].astype(str).tolist()]) if len(df) else len(str(col_name))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


tab1, tab2, tab3, tab4 = st.tabs(
    ["🚀 Paleidimas", "📤 Įkelti failą", "✅ Patvirtinimas", "📈 Suvestinė"]
)

# ------------------------------------------------------------
# 1. PALEIDIMAS
# ------------------------------------------------------------
with tab1:
    st.subheader("Paleisti apdorojimą dabar")
    st.write(
        "Automatika ir taip veikia pati kas 10 minučių. "
        "Šis mygtukas leidžia paleisti apdorojimą iškart, nelaukiant."
    )
    if st.button("▶️ Apdoroti dabar", type="primary"):
        with st.spinner("Siunčiama komanda į n8n..."):
            try:
                resp = requests.post(
                    f"{N8N_BASE_URL}/webhook/apdoroti-dabar",
                    headers=headers(),
                    timeout=REQUEST_TIMEOUT,
                )
                if resp.status_code in (200, 201):
                    st.success("Apdorojimas paleistas! Rezultatų ieškok skiltyje „Suvestinė“ po kelių minučių.")
                    fetch_data.clear()
                else:
                    st.error(f"n8n atsakė klaida (kodas {resp.status_code}). Patikrink Webhook mazgo nustatymus.")
            except requests.exceptions.RequestException as e:
                st.error(f"Nepavyko pasiekti n8n. Ar veikia tunelis ir n8n konteineris? Klaida: {e}")

# ------------------------------------------------------------
# 2. FAILŲ ĮKĖLIMAS
# ------------------------------------------------------------
with tab2:
    st.subheader("Įkelti naują dokumentą")
    st.write("Įkelk sutarties priedą, sąskaitą ar atliktų darbų aktą (PDF arba Word).")

    uploaded = st.file_uploader("Pasirink failą", type=["pdf", "docx", "doc"])
    if uploaded is not None:
        if st.button("⬆️ Įkelti į apdorojimo eilę"):
            with st.spinner("Keliama..."):
                try:
                    upload_headers = dict(headers())
                    upload_headers["Content-Type"] = "application/octet-stream"
                    upload_headers["X-Filename"] = quote(uploaded.name)
                    resp = requests.post(
                        f"{N8N_BASE_URL}/webhook/ikelti-faila",
                        headers=upload_headers,
                        data=uploaded.getvalue(),
                        timeout=60,
                    )
                    resp.raise_for_status()
                    result = resp.json()
                    status = result.get("status")
                    if status == "exists":
                        st.warning(f"„{uploaded.name}“ jau yra aplanke arba jau apdorotas anksčiau.")
                    elif status == "ok":
                        st.success(
                            f"„{uploaded.name}“ įkeltas. Bus apdorotas per artimiausius 10 min "
                            f"(arba paspausk „Apdoroti dabar“ skiltyje „Paleidimas“)."
                        )
                    else:
                        st.error(f"n8n grąžino netikėtą atsakymą: {result}")
                except requests.exceptions.RequestException as e:
                    st.error(f"Nepavyko įkelti failo. Klaida: {e}")

# ------------------------------------------------------------
# 3. PATVIRTINIMAS
# ------------------------------------------------------------
with tab3:
    st.subheader("Laukiančios patvirtinimo eilutės")

    try:
        df = fetch_data()
    except requests.exceptions.RequestException as e:
        df = None
        st.error(f"Nepavyko gauti duomenų iš n8n. Klaida: {e}")

    if df is not None:
        if df.empty:
            st.info("Duomenų dar nėra.")
        else:
            df_data = df[df["klientas"] != "VISO:"].copy()

            if "busena" not in df_data.columns:
                st.error("Stulpelio „busena“ duomenyse nėra — patikrink n8n endpoint'ą.")
            else:
                pending = df_data[df_data["busena"] == "Laukia"]

                if len(pending) == 0:
                    st.info("Nėra laukiančių patvirtinimo eilučių. ✅")
                else:
                    st.write(f"Iš viso laukia patvirtinimo: **{len(pending)}**")

                    edited = st.data_editor(
                        pending,
                        column_config={
                            "busena": st.column_config.SelectboxColumn(
                                "busena", options=ALLOWED_STATUSES, required=True
                            )
                        },
                        disabled=[c for c in pending.columns if c != "busena"],
                        use_container_width=True,
                        hide_index=True,
                        key="review_table",
                    )

                    if st.button("💾 Išsaugoti pakeitimus"):
                        changed = edited[edited["busena"] != pending["busena"]]
                        if changed.empty:
                            st.info("Nėra pakeitimų.")
                        else:
                            updates = [
                                {"saltinio_failas": row["saltinio_failas"], "busena": row["busena"]}
                                for _, row in changed.iterrows()
                            ]
                            try:
                                post_status_updates(updates)
                                fetch_data.clear()
                                st.success("Išsaugota!")
                                st.rerun()
                            except requests.exceptions.RequestException as e:
                                st.error(f"Nepavyko išsaugoti. Klaida: {e}")

# ------------------------------------------------------------
# 4. SUVESTINĖ
# ------------------------------------------------------------
with tab4:
    st.subheader("Būsenos suvestinė")

    try:
        df = fetch_data()
    except requests.exceptions.RequestException as e:
        df = None
        st.error(f"Nepavyko gauti duomenų iš n8n. Klaida: {e}")

    if df is not None and not df.empty:
        df_data = df[df["klientas"] != "VISO:"].copy()

        laukia = int((df_data["busena"] == "Laukia").sum()) if "busena" in df_data.columns else 0
        patvirtinta = int((df_data["busena"] == "Patvirtinta").sum()) if "busena" in df_data.columns else 0

        col1, col2, col3 = st.columns(3)
        col1.metric("Iš viso eilučių", len(df_data))
        col2.metric("Laukia", laukia)
        col3.metric("Patvirtinta", patvirtinta)

        viso_row = df[df["klientas"] == "VISO:"]
        if len(viso_row) > 0 and "suma_su_pvm" in viso_row.columns:
            st.metric("Bendra suma (su PVM)", f"{float(viso_row.iloc[0]['suma_su_pvm']):.2f} €")

        st.caption(f"Atnaujinta: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

        col_a, col_b = st.columns(2)
        with col_a:
            if st.button("🔄 Atnaujinti"):
                fetch_data.clear()
                st.rerun()
        with col_b:
            st.download_button(
                label="⬇️ Atsisiųsti Excel failą",
                data=to_excel_bytes(df),
                file_name=f"sumu_suvestine_{datetime.now().strftime('%Y-%m-%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
    elif df is not None:
        st.warning("Dar nėra duomenų.")
